from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None

try:
    import gi

    gi.require_version("Gtk", "3.0")
    from gi.repository import Gtk  # noqa: F401
except Exception:  # pragma: no cover
    Gtk = None

from .config import MeasurementMode


@dataclass(frozen=True)
class RecordingColumn:
    key: str
    device_name: str
    quantity: str
    label: str
    unit: str

    @property
    def heading(self) -> str:
        return f"{self.label} ({self.unit}) - {self.device_name}"


@dataclass(frozen=True)
class RecordingTarget:
    path: Path
    format: str

    @property
    def exists(self) -> bool:
        return self.path.exists()


@dataclass(frozen=True)
class RecordingSpool:
    target: RecordingTarget
    active_path: Path
    temp_path: Optional[Path]


def canonical_unit_for_mode(mode: MeasurementMode) -> str:
    if mode == "voltage":
        return "V"
    if mode == "current":
        return "A"
    if mode == "resistance":
        return "Ohm"
    if mode == "temperature":
        return "C"
    if mode == "frequency":
        return "Hz"
    return "F"


def label_for_mode(mode: MeasurementMode) -> str:
    if mode == "voltage":
        return "Voltage"
    if mode == "current":
        return "Current"
    if mode == "resistance":
        return "Resistance"
    if mode == "temperature":
        return "Temperature"
    if mode == "frequency":
        return "Frequency"
    return "Capacitance"


def write_rows(target: RecordingTarget, headers: list[str], rows: list[list[object]]) -> None:
    target.path.parent.mkdir(parents=True, exist_ok=True)
    if target.format == "csv":
        _write_csv(target.path, headers, rows)
        return
    _write_xlsx(target.path, headers, rows)


def prepare_recording_spool(target: RecordingTarget, headers: list[str]) -> RecordingSpool:
    target.path.parent.mkdir(parents=True, exist_ok=True)
    if target.format == "csv":
        _write_csv(target.path, headers, [])
        return RecordingSpool(target=target, active_path=target.path, temp_path=None)

    temp_path = target.path.with_name(f".{target.path.stem}.recording.csv")
    _write_csv(temp_path, headers, [])
    return RecordingSpool(target=target, active_path=temp_path, temp_path=temp_path)


def append_recording_rows(spool: RecordingSpool, rows: list[list[object]]) -> None:
    if not rows:
        return
    with spool.active_path.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)


def finalize_recording_spool(spool: RecordingSpool) -> None:
    if spool.target.format == "csv":
        return
    assert spool.temp_path is not None
    _csv_to_xlsx(spool.temp_path, spool.target.path)
    spool.temp_path.unlink(missing_ok=True)


def cleanup_recording_spool(spool: Optional[RecordingSpool]) -> None:
    if spool is None:
        return
    if spool.temp_path is not None:
        spool.temp_path.unlink(missing_ok=True)


def ensure_format_supported(format_name: str) -> None:
    if format_name == "xlsx" and Workbook is None:
        raise RuntimeError("Excel export requires the 'openpyxl' package")


def _write_csv(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def _write_xlsx(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    if Workbook is None:
        raise RuntimeError("Excel export requires the 'openpyxl' package")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _csv_to_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    if Workbook is None:
        raise RuntimeError("Excel export requires the 'openpyxl' package")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title="Measurements")
    with csv_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        for row in reader:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        default_sheet = workbook["Sheet"]
        workbook.remove(default_sheet)
    workbook.save(xlsx_path)


def parse_recording_target(raw_path: str) -> RecordingTarget:
    return _parse_recording_target(raw_path, fallback_format="csv")


def parse_recording_target_from_dialog(raw_path: str, selected_format: str) -> RecordingTarget:
    return _parse_recording_target(raw_path, fallback_format=selected_format)


def _parse_recording_target(raw_path: str, fallback_format: str) -> RecordingTarget:
    normalized = raw_path.strip()
    if not normalized:
        raise RuntimeError("Enter a recording file path or use Choose File")

    expanded = Path(normalized).expanduser()
    if not expanded.is_absolute():
        expanded = Path.cwd() / expanded

    suffix = expanded.suffix.lower()
    if suffix == ".xlsx":
        format_name = "xlsx"
    elif suffix == ".csv":
        format_name = "csv"
    else:
        format_name = fallback_format
    if not suffix:
        expanded = expanded.with_suffix(".xlsx" if format_name == "xlsx" else ".csv")
    elif suffix and suffix not in {".csv", ".xlsx"}:
        format_name = fallback_format

    ensure_format_supported(format_name)
    return RecordingTarget(path=expanded, format=format_name)


def choose_recording_path(current_path: Optional[str]) -> RecordingTarget:
    if Gtk is None:
        raise RuntimeError("Desktop file chooser is unavailable in this launch environment")

    current = _default_dialog_filename(current_path)
    helper = textwrap.dedent(
        """
        import json
        import sys
        from pathlib import Path

        import gi
        gi.require_version("Gtk", "3.0")
        from gi.repository import Gtk

        current = Path(sys.argv[1]).expanduser()
        dialog = Gtk.FileChooserDialog(
            "Choose recording file",
            None,
            Gtk.FileChooserAction.SAVE,
            (
                Gtk.STOCK_CANCEL,
                Gtk.ResponseType.CANCEL,
                Gtk.STOCK_SAVE,
                Gtk.ResponseType.ACCEPT,
            ),
        )
        dialog.set_do_overwrite_confirmation(False)
        dialog.set_modal(True)

        if current.exists():
            dialog.set_filename(str(current))
        else:
            parent = current.parent if str(current.parent) else Path.home()
            if parent.exists():
                dialog.set_current_folder(str(parent))
            dialog.set_current_name(current.name)

        csv_filter = Gtk.FileFilter()
        csv_filter.set_name("CSV files")
        csv_filter.add_pattern("*.csv")
        dialog.add_filter(csv_filter)

        xlsx_filter = Gtk.FileFilter()
        xlsx_filter.set_name("Excel files")
        xlsx_filter.add_pattern("*.xlsx")
        dialog.add_filter(xlsx_filter)

        all_filter = Gtk.FileFilter()
        all_filter.set_name("All files")
        all_filter.add_pattern("*")
        dialog.add_filter(all_filter)

        if current.suffix.lower() == ".xlsx":
            dialog.set_filter(xlsx_filter)
        else:
            dialog.set_filter(csv_filter)

        response = dialog.run()
        if response != Gtk.ResponseType.ACCEPT:
            dialog.destroy()
            sys.exit(2)

        selected_path = dialog.get_filename()
        current_folder = dialog.get_current_folder()
        current_name = dialog.get_current_name()
        if not selected_path and current_folder and current_name:
            selected_path = str(Path(current_folder) / current_name)
        elif not selected_path and current_name:
            selected_path = current_name

        result = {
            "path": selected_path,
            "current_folder": current_folder,
            "current_name": current_name,
            "filter_name": None if dialog.get_filter() is None else dialog.get_filter().get_name(),
        }
        dialog.destroy()
        print(json.dumps(result))
        """
    )
    completed = subprocess.run(
        [sys.executable, "-c", helper, current],
        check=False,
        capture_output=True,
        text=True,
    )

    if completed.returncode == 2:
        raise RuntimeError("No file selected")
    if completed.returncode != 0:
        detail = completed.stderr.strip() or "Unknown chooser error"
        raise RuntimeError(f"Choose File failed: {detail}")

    output = completed.stdout.strip()
    if not output:
        raise RuntimeError("No file selected")
    payload = json.loads(output)
    selected = payload.get("path", "").strip()
    current_folder = (payload.get("current_folder") or "").strip()
    current_name = (payload.get("current_name") or "").strip()
    selected_filter = payload.get("filter_name")
    if not selected and current_name and current_folder:
        selected = str(Path(current_folder) / current_name)
    elif not selected and current_name:
        selected = current_name
    if not selected:
        raise RuntimeError(
            "Choose File failed: dialog did not return a filename "
            f"(filter={selected_filter or 'unknown'})"
        )

    suffix = Path(selected).suffix.lower()
    if suffix in {".csv", ".xlsx"}:
        return parse_recording_target_from_dialog(
            selected,
            "xlsx" if suffix == ".xlsx" else "csv",
        )

    fallback_format = _format_from_filter_name(selected_filter)
    return parse_recording_target_from_dialog(selected, fallback_format)


def _default_dialog_filename(current_path: Optional[str]) -> str:
    if current_path and current_path.strip():
        return str(Path(current_path.strip()).expanduser())
    return str(Path.home() / "measurements.csv")


def _format_from_filter_name(filter_name: Optional[str]) -> str:
    if filter_name and "excel" in filter_name.lower():
        return "xlsx"
    return "csv"


def recording_file_dialog_available() -> tuple[bool, Optional[str]]:
    if os.getenv("SURFACE_CONDUCTIVITY_DESKTOP") == "1":
        return True, None

    if Gtk is None:
        return False, "Desktop file chooser is unavailable; enter the file path manually."

    test_script = textwrap.dedent(
        """
        import gi
        gi.require_version("Gtk", "3.0")
        from gi.repository import Gtk
        Gtk.init_check()
        """
    )
    completed = subprocess.run(
        [sys.executable, "-c", test_script],
        check=False,
        capture_output=True,
        text=True,
    )
    if completed.returncode != 0:
        return False, "Desktop file chooser cannot open from this app session; enter the file path manually."
    return True, None
